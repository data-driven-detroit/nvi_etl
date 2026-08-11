"""Generate per-CDO Excel workbooks from the primary_survey_cdo CSV.

Each CDO gets its own .xlsx file with:
  - A data sheet comparing CDO responses to citywide responses
  - A map sheet showing the CDO boundary within Detroit

Reads the CSV produced by primary_survey_cdo and CDO geometries from the
database.
"""

import io
from pathlib import Path

import folium
import geopandas as gpd
import pandas as pd
from folium import DivIcon
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
from sqlalchemy import Engine

from nvi_etl.geo import pull_cdo_boundaries, pull_city_boundary
from nvi_etl.registry import task, TaskResult
from nvi_etl.tasks.primary_survey import SURVEY_YEAR

INPUT_DIR = Path(__file__).resolve().parent.parent / "survey" / "output"
OUTPUT_DIR = INPUT_DIR / "cdo_workbooks"

HEADER_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
HEADER_FONT = Font(name="IBM Plex Sans", bold=True, size=11)
BODY_FONT = Font(name="IBM Plex Sans", size=10)
ZEBRA_FILL = PatternFill(start_color="F7F7F7", end_color="F7F7F7", fill_type="solid")
WRAP = Alignment(wrap_text=True, vertical="top")


# ---------------------------------------------------------------------------
# Map generation
# ---------------------------------------------------------------------------

def _generate_cdo_map(cdo_geom, city_geom, cdo_name):
    """Render a CDO boundary on a Detroit base map and return PNG bytes."""
    detroit_json = gpd.GeoSeries(city_geom.to_crs(4326)["geometry"]).simplify(0.001).to_json()
    cdo_json = gpd.GeoSeries(cdo_geom.to_crs(4326)["geometry"]).simplify(0.001).to_json()
    cdo_centroid = gpd.GeoSeries(cdo_geom.to_crs(4326)["geometry"]).simplify(0.001).centroid.iloc[0]

    cdo_bounds = gpd.GeoSeries(cdo_geom.to_crs(4326)["geometry"]).total_bounds
    sw = [cdo_bounds[1], cdo_bounds[0]]
    ne = [cdo_bounds[3], cdo_bounds[2]]

    m = folium.Map(tiles="CartoDB positron")
    m.fit_bounds([sw, ne], padding=[20, 20])

    folium.GeoJson(
        detroit_json,
        style_function=lambda x: {
            "fillColor": "none", "color": "black", "weight": 2,
        },
    ).add_to(m)

    folium.GeoJson(
        cdo_json,
        style_function=lambda x: {
            "fillColor": "#87AF3F", "color": "#87AF3F",
            "weight": 2, "fillOpacity": 0.6,
        },
    ).add_to(m)

    img_data = m._to_png(20)
    img = PILImage.open(io.BytesIO(img_data))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Workbook creation
# ---------------------------------------------------------------------------

def _write_data_sheet(ws, cdo_data, citywide_data, cdo_name):
    """Write the side-by-side CDO vs. citywide comparison sheet."""
    headers = [
        "Topic", "Question", "Answer",
        f"{cdo_name}\nCount", f"{cdo_name}\nUniverse", f"{cdo_name}\n%",
        "Citywide\nCount", "Citywide\nUniverse", "Citywide\n%",
    ]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP

    # Build a lookup for citywide rows keyed on (topic, question, answer, value_type)
    cw_lookup = {}
    for _, row in citywide_data.iterrows():
        key = (row["topic_text"], row["question_text"], row["answer"], row["value_type"])
        cw_lookup[key] = row

    row_num = 2
    for _, row in cdo_data.iterrows():
        key = (row["topic_text"], row["question_text"], row["answer"], row["value_type"])
        cw = cw_lookup.get(key)

        values = [
            row["topic_text"], row["question_text"], row["answer"],
            row["count"], row["universe"], row["percentage"],
            cw["count"] if cw is not None else None,
            cw["universe"] if cw is not None else None,
            cw["percentage"] if cw is not None else None,
        ]
        fill = ZEBRA_FILL if row_num % 2 == 0 else None
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = BODY_FONT
            if fill:
                cell.fill = fill

        row_num += 1

    # Column widths
    col_widths = [25, 40, 25, 12, 12, 12, 12, 12, 12]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_map_sheet(ws, map_bytes, cdo_name):
    """Add a CDO map image to its own sheet."""
    ws.column_dimensions["A"].width = 100
    ws["A1"] = f"{cdo_name} — Service Area Boundary"
    ws["A1"].font = Font(name="IBM Plex Sans", bold=True, size=14)

    img = XlImage(map_bytes)
    img.width = 1000
    img.height = 650
    ws.add_image(img, "A3")


def create_cdo_workbook(cdo_name, cdo_data, citywide_data, cdo_geom, city_geom):
    """Build a single CDO workbook and return the Workbook object."""
    wb = Workbook()

    # Data sheet
    ws_data = wb.active
    ws_data.title = "Survey Data"
    _write_data_sheet(ws_data, cdo_data, citywide_data, cdo_name)

    # Map sheet
    ws_map = wb.create_sheet("Boundary Map")
    map_bytes = _generate_cdo_map(cdo_geom, city_geom, cdo_name)
    _write_map_sheet(ws_map, map_bytes, cdo_name)

    return wb


# ---------------------------------------------------------------------------
# Task entry point
# ---------------------------------------------------------------------------

@task("cdo_workbooks", phase=2, description="Per-CDO Excel workbooks with maps and citywide comparison")
def run(source: Engine, target: Engine) -> TaskResult:
    import logging
    logger = logging.getLogger("nvi_etl")

    csv_path = INPUT_DIR / f"primary_survey_cdo_{SURVEY_YEAR}.csv"
    if not csv_path.exists():
        raise FileNotFoundError(
            f"{csv_path} not found — run primary_survey_cdo first"
        )

    logger.info(f"Reading CDO survey data from {csv_path}")
    data = pd.read_csv(csv_path, dtype={"count": str, "universe": str, "percentage": str})

    citywide_data = data[data["organization_name"] == "citywide"]
    cdo_data = data[data["organization_name"] != "citywide"]
    cdo_names = cdo_data["organization_name"].dropna().unique()

    # Pull geometries for maps
    logger.info("Pulling geometries for map generation")
    cdo_boundaries = pull_cdo_boundaries(source)
    city_boundary = pull_city_boundary(source)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbooks_created = 0

    for cdo_name in sorted(cdo_names):
        logger.info(f"Creating workbook for {cdo_name}")
        cdo_rows = cdo_data[cdo_data["organization_name"] == cdo_name]
        cdo_geom = cdo_boundaries[cdo_boundaries["organization_name"] == cdo_name]

        if cdo_geom.empty:
            logger.warning(f"No geometry found for {cdo_name}, skipping")
            continue

        wb = create_cdo_workbook(
            cdo_name, cdo_rows, citywide_data, cdo_geom, city_boundary
        )

        safe_name = cdo_name.replace("/", "-").replace("\\", "-")
        wb.save(OUTPUT_DIR / f"{safe_name}.xlsx")
        workbooks_created += 1

    logger.info(f"Created {workbooks_created} CDO workbooks in {OUTPUT_DIR}")

    return TaskResult(task_name="cdo_workbooks", rows_inserted=workbooks_created, success=True)
